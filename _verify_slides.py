from pptx import Presentation
prs = Presentation('CMF_Unlearning_Analysis.pptx')
print(f'Slides: {len(prs.slides)}')
print(f'Canvas: {prs.slide_width.inches:.2f}" x {prs.slide_height.inches:.2f}"')
for i, sl in enumerate(prs.slides):
    shapes = len(sl.shapes)
    texts = [sh.text_frame.text[:50].strip() for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()][:2]
    print(f'  Slide {i}: {shapes} shapes | {texts}')
